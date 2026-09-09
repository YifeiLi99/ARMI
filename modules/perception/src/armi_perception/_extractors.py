"""Deterministic extraction and media validation for external message parts."""

from __future__ import annotations

import csv
import io
import json
import warnings
from dataclasses import dataclass
from pathlib import Path
from zipfile import BadZipFile, ZipFile

from armi_interaction.api import (
    ExternalMessagePartKind,
    ExternalMessageViolation,
    ExternalVisualRole,
)
from docx import Document
from docx.opc.exceptions import PackageNotFoundError as DocxPackageNotFoundError
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from PIL import Image, UnidentifiedImageError
from pptx import Presentation
from pptx.exc import PackageNotFoundError as PptxPackageNotFoundError

from .api import ExternalMediaContent

_MAX_PROJECTION_BYTES = 256 * 1024
_TRUNCATED = "\n[内容已按 ARMI 单条认知材料上限截断]"
_MAX_IMAGE_PIXELS = 36_000_000
_MAX_IMAGE_FRAMES = 256
_MAX_IMAGE_TOTAL_PIXELS = 72_000_000
_MAX_RENDERED_PNG_BYTES = 10 * 1024 * 1024
_MAX_RENDERED_PNG_TOTAL_BYTES = 25 * 1024 * 1024
_MAX_ZIP_MEMBERS = 2_048
_MAX_ZIP_MEMBER_BYTES = 32 * 1024 * 1024
_MAX_ZIP_TOTAL_BYTES = 64 * 1024 * 1024
_MAX_XML_MEMBER_BYTES = 8 * 1024 * 1024
_MAX_XML_TOTAL_BYTES = 32 * 1024 * 1024
_MAX_STRUCTURE_UNITS = 100_000
_GENERIC_STICKER_SUMMARIES = frozenset(
    {
        "图片",
        "表情",
        "表情包",
        "动画表情",
        "商城表情",
        "[图片]",
        "[表情]",
        "[动画表情]",
        "[商城表情]",
    }
)


@dataclass(frozen=True, slots=True)
class ExtractedExternalContent:
    media_type: str
    text: str | None
    requires_provider: bool
    pixel_width: int | None = None
    pixel_height: int | None = None
    frame_count: int | None = None
    visual_inputs: tuple[ExternalMediaContent, ...] = ()


def extract_external_content(
    *,
    kind: ExternalMessagePartKind,
    content: bytes,
    file_name: str,
    visual_role: ExternalVisualRole | None = None,
    source_summary: str | None = None,
) -> ExtractedExternalContent:
    if kind is ExternalMessagePartKind.IMAGE:
        if type(visual_role) is not ExternalVisualRole:
            raise ExternalMessageViolation("EXTERNAL-MESSAGE-MEDIA-TYPE")
        return _extract_image(
            content,
            file_name=file_name,
            visual_role=visual_role,
            source_summary=source_summary,
        )
    if kind is ExternalMessagePartKind.AUDIO:
        if not (content.startswith(b"ID3") or _looks_like_mp3_frame(content)):
            raise ExternalMessageViolation("EXTERNAL-MESSAGE-MEDIA-TYPE")
        return ExtractedExternalContent("audio/mpeg", None, True)
    if kind is ExternalMessagePartKind.VIDEO:
        if len(content) < 12 or content[4:8] != b"ftyp":
            raise ExternalMessageViolation("EXTERNAL-MESSAGE-MEDIA-TYPE")
        return ExtractedExternalContent("video/mp4", None, True)
    if kind is not ExternalMessagePartKind.FILE:
        raise ExternalMessageViolation("EXTERNAL-MESSAGE-MEDIA-TYPE")
    suffix = Path(file_name).suffix.lower()
    if content.startswith(b"%PDF-"):
        return ExtractedExternalContent("application/pdf", None, True)
    if suffix in {".txt", ".md", ".markdown", ".log", ".json", ".csv"} or suffix in {
        ".py",
        ".js",
        ".ts",
        ".tsx",
        ".jsx",
        ".java",
        ".c",
        ".h",
        ".cpp",
        ".cs",
        ".go",
        ".rs",
        ".sql",
        ".toml",
        ".yaml",
        ".yml",
        ".xml",
        ".html",
        ".css",
    }:
        text = _decode_text(content)
        if suffix == ".json":
            try:
                text = json.dumps(json.loads(text), ensure_ascii=False, indent=2)
            except json.JSONDecodeError:
                raise ExternalMessageViolation(
                    "EXTERNAL-MESSAGE-FILE-INVALID"
                ) from None
        elif suffix == ".csv":
            text = "\n".join(" | ".join(row) for row in csv.reader(io.StringIO(text)))
        return ExtractedExternalContent("text/plain", _bounded(text), False)
    try:
        _preflight_ooxml(content)
        office_kind = _office_kind(content)
        text = (
            _docx_text(content)
            if office_kind == "docx"
            else _pptx_text(content)
            if office_kind == "pptx"
            else _xlsx_text(content)
            if office_kind == "xlsx"
            else None
        )
    except (
        BadZipFile,
        DocxPackageNotFoundError,
        InvalidFileException,
        KeyError,
        OSError,
        PptxPackageNotFoundError,
        SyntaxError,
        ValueError,
    ):
        raise ExternalMessageViolation("EXTERNAL-MESSAGE-FILE-INVALID") from None
    if office_kind == "docx":
        return ExtractedExternalContent(
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            _bounded(text or ""),
            False,
        )
    if office_kind == "pptx":
        return ExtractedExternalContent(
            "application/vnd.openxmlformats-officedocument.presentationml.presentation",
            _bounded(text or ""),
            False,
        )
    if office_kind == "xlsx":
        return ExtractedExternalContent(
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            _bounded(text or ""),
            False,
        )
    raise ExternalMessageViolation("EXTERNAL-MESSAGE-FILE-UNSUPPORTED")


def _office_kind(content: bytes) -> str | None:
    if not content.startswith(b"PK"):
        return None
    with ZipFile(io.BytesIO(content)) as archive:
        names = frozenset(archive.namelist())
    if "word/document.xml" in names:
        return "docx"
    if "ppt/presentation.xml" in names:
        return "pptx"
    if "xl/workbook.xml" in names:
        return "xlsx"
    return None


def _preflight_ooxml(content: bytes) -> None:
    if not content.startswith(b"PK"):
        return
    total = 0
    xml_total = 0
    with ZipFile(io.BytesIO(content)) as archive:
        members = archive.infolist()
        if len(members) > _MAX_ZIP_MEMBERS:
            raise ExternalMessageViolation("EXTERNAL-MESSAGE-RESOURCE-LIMIT")
        for member in members:
            if member.is_dir():
                continue
            size = member.file_size
            total += size
            normalized = member.filename.replace("\\", "/")
            if (
                size > _MAX_ZIP_MEMBER_BYTES
                or total > _MAX_ZIP_TOTAL_BYTES
                or normalized.startswith("/")
                or ".." in normalized.split("/")
                or size > max(1024 * 1024, member.compress_size * 100)
            ):
                raise ExternalMessageViolation("EXTERNAL-MESSAGE-RESOURCE-LIMIT")
            if normalized.casefold().endswith((".xml", ".rels")):
                xml_total += size
                if size > _MAX_XML_MEMBER_BYTES or xml_total > _MAX_XML_TOTAL_BYTES:
                    raise ExternalMessageViolation("EXTERNAL-MESSAGE-RESOURCE-LIMIT")


def _decode_text(content: bytes) -> str:
    for encoding in ("utf-8-sig", "gb18030"):
        try:
            return content.decode(encoding, errors="strict")
        except UnicodeDecodeError:
            continue
    raise ExternalMessageViolation("EXTERNAL-MESSAGE-FILE-INVALID")


def _bounded(value: str) -> str:
    if not value.strip():
        raise ExternalMessageViolation("EXTERNAL-MESSAGE-FILE-EMPTY")
    encoded = value.encode("utf-8", errors="strict")
    if len(encoded) <= _MAX_PROJECTION_BYTES:
        return value
    limit = _MAX_PROJECTION_BYTES - len(_TRUNCATED.encode("utf-8"))
    clipped = encoded[:limit]
    while True:
        try:
            return clipped.decode("utf-8", errors="strict") + _TRUNCATED
        except UnicodeDecodeError:
            clipped = clipped[:-1]


def _image_media_type(content: bytes) -> str:
    signatures = (
        (b"\xff\xd8\xff", "image/jpeg"),
        (b"\x89PNG\r\n\x1a\n", "image/png"),
        (b"GIF87a", "image/gif"),
        (b"GIF89a", "image/gif"),
        (b"BM", "image/bmp"),
    )
    for signature, media_type in signatures:
        if content.startswith(signature):
            return media_type
    if len(content) >= 12 and content.startswith(b"RIFF") and content[8:12] == b"WEBP":
        return "image/webp"
    raise ExternalMessageViolation("EXTERNAL-MESSAGE-MEDIA-TYPE")


def _extract_image(
    content: bytes,
    *,
    file_name: str,
    visual_role: ExternalVisualRole,
    source_summary: str | None,
) -> ExtractedExternalContent:
    media_type = _image_media_type(content)
    try:
        with warnings.catch_warnings():
            warnings.simplefilter("error", Image.DecompressionBombWarning)
            with Image.open(io.BytesIO(content)) as image:
                width, height = image.size
                frame_count = int(getattr(image, "n_frames", 1))
                if (
                    width <= 0
                    or height <= 0
                    or width * height > _MAX_IMAGE_PIXELS
                    or frame_count <= 0
                    or frame_count > _MAX_IMAGE_FRAMES
                ):
                    raise ExternalMessageViolation("EXTERNAL-MESSAGE-RESOURCE-LIMIT")
                total_pixels = 0
                for index in range(frame_count):
                    image.seek(index)
                    frame_width, frame_height = image.size
                    total_pixels += frame_width * frame_height
                    if total_pixels > _MAX_IMAGE_TOTAL_PIXELS:
                        raise ExternalMessageViolation(
                            "EXTERNAL-MESSAGE-RESOURCE-LIMIT"
                        )
                image.verify()
            inputs = _visual_inputs(
                content,
                file_name=file_name,
                media_type=media_type,
                frame_count=frame_count,
            )
    except ExternalMessageViolation:
        raise
    except Image.DecompressionBombError, Image.DecompressionBombWarning:
        raise ExternalMessageViolation("EXTERNAL-MESSAGE-IMAGE-DIMENSIONS") from None
    except OSError, SyntaxError, UnidentifiedImageError, ValueError:
        raise ExternalMessageViolation("EXTERNAL-MESSAGE-FILE-INVALID") from None
    local_text = None
    if visual_role is ExternalVisualRole.STICKER and _meaningful_sticker_summary(
        source_summary
    ):
        local_text = f"QQ 提供的商城表情摘要: {source_summary}"
    return ExtractedExternalContent(
        media_type,
        local_text,
        local_text is None,
        width,
        height,
        frame_count,
        inputs,
    )


def _visual_inputs(
    content: bytes, *, file_name: str, media_type: str, frame_count: int
) -> tuple[ExternalMediaContent, ...]:
    if frame_count == 1 and media_type in {"image/jpeg", "image/png"}:
        return (ExternalMediaContent(content, file_name, media_type),)
    indexes = _frame_indexes(frame_count)
    values: list[ExternalMediaContent] = []
    total_output_bytes = 0
    with Image.open(io.BytesIO(content)) as image:
        for ordinal, index in enumerate(indexes, start=1):
            image.seek(index)
            frame = image.convert("RGBA")
            output = io.BytesIO()
            frame.save(output, format="PNG")
            rendered = output.getvalue()
            total_output_bytes += len(rendered)
            if (
                len(rendered) > _MAX_RENDERED_PNG_BYTES
                or total_output_bytes > _MAX_RENDERED_PNG_TOTAL_BYTES
            ):
                raise ExternalMessageViolation("EXTERNAL-MESSAGE-RESOURCE-LIMIT")
            values.append(
                ExternalMediaContent(
                    rendered,
                    f"{Path(file_name).stem}-frame-{ordinal}.png",
                    "image/png",
                )
            )
    return tuple(values)


def _frame_indexes(frame_count: int) -> tuple[int, ...]:
    if frame_count <= 4:
        return tuple(range(frame_count))
    return tuple(round(index * (frame_count - 1) / 3) for index in range(4))


def _meaningful_sticker_summary(value: str | None) -> bool:
    if value is None:
        return False
    normalized = value.strip()
    return bool(normalized) and normalized.casefold() not in {
        item.casefold() for item in _GENERIC_STICKER_SUMMARIES
    }


def _looks_like_mp3_frame(content: bytes) -> bool:
    return len(content) >= 2 and content[0] == 0xFF and content[1] & 0xE0 == 0xE0


def _docx_text(content: bytes) -> str:
    document = Document(io.BytesIO(content))
    collector = _ProjectionCollector()
    for paragraph in document.paragraphs:
        if not collector.add(paragraph.text):
            return collector.finish()
    for table_index, table in enumerate(document.tables, start=1):
        if not collector.add(f"[表格 {table_index}]"):
            return collector.finish()
        for row in table.rows:
            if not collector.add(
                " | ".join(cell.text for cell in row.cells), units=len(row.cells)
            ):
                return collector.finish()
    return collector.finish()


def _pptx_text(content: bytes) -> str:
    presentation = Presentation(io.BytesIO(content))
    collector = _ProjectionCollector()
    for index, slide in enumerate(presentation.slides, start=1):
        values: list[str] = []
        for shape in slide.shapes:
            collector.count_unit()
            text = getattr(shape, "text", None)
            if isinstance(text, str) and text:
                values.append(text)
        if not collector.add(f"[幻灯片 {index}]\n" + "\n".join(values)):
            return collector.finish()
    return collector.finish(separator="\n\n")


def _xlsx_text(content: bytes) -> str:
    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=False)
    collector = _ProjectionCollector()
    try:
        for sheet in workbook.worksheets:
            if not collector.add(f"[工作表 {sheet.title}]"):
                return collector.finish()
            for row in sheet.iter_rows(values_only=True):
                if not collector.add(
                    " | ".join("" if value is None else str(value) for value in row),
                    units=len(row),
                ):
                    return collector.finish()
    finally:
        workbook.close()
    return collector.finish()


class _ProjectionCollector:
    __slots__ = ("_chunks", "_size", "_truncated", "_units")

    def __init__(self) -> None:
        self._chunks: list[str] = []
        self._size = 0
        self._truncated = False
        self._units = 0

    def count_unit(self, count: int = 1) -> None:
        self._units += count
        if self._units > _MAX_STRUCTURE_UNITS:
            raise ExternalMessageViolation("EXTERNAL-MESSAGE-RESOURCE-LIMIT")

    def add(self, value: str, *, units: int = 1) -> bool:
        self.count_unit(units)
        if not value:
            return True
        separator_bytes = 1 if self._chunks else 0
        encoded = value.encode("utf-8", errors="strict")
        room = _MAX_PROJECTION_BYTES - len(_TRUNCATED.encode("utf-8"))
        if self._size + separator_bytes + len(encoded) <= room:
            self._chunks.append(value)
            self._size += separator_bytes + len(encoded)
            return True
        self._truncated = True
        return False

    def finish(self, *, separator: str = "\n") -> str:
        value = separator.join(self._chunks)
        return value + _TRUNCATED if self._truncated else value


__all__ = ("ExtractedExternalContent", "extract_external_content")
